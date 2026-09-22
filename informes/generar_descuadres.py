import json
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

d = json.load(open('/tmp/claude-0/datos.json'))
d.sort(key=lambda x: (x['cliente'].upper(), x['fecha']))

FUENTE = 'Arial'
AZUL   = 'FF1F3864'
GRIS   = 'FFF2F2F2'
CAB    = PatternFill('solid', fgColor='FF1F3864')
FCLI   = PatternFill('solid', fgColor='FFDDEBF7')
FSUB   = PatternFill('solid', fgColor='FFF2F2F2')
FEDIT  = PatternFill('solid', fgColor='FFFFFF99')
fino   = Side(style='thin', color='FFBFBFBF')
BORDE  = Border(left=fino, right=fino, top=fino, bottom=fino)
EUR    = '#,##0.00 "€";[Red](#,##0.00 "€");-'
FECHA  = 'DD/MM/YYYY'

ESTFAC = {0:'0 Pendiente', 1:'1 Cobro parcial', 2:'2 Cobrada', 3:'3 Devuelta', 4:'4 Anulada'}

def diagnostico(x):
    col, cob, tot = x['importe_cobrado'], x['cobros_anotados'], x['total']
    if col > tot + 0.01:
        return 'La columna pasa del total de la factura: parece contado dos veces'
    if col == 0 and cob > 0:
        return 'La columna esta a cero pero si hay cobros anotados'
    if abs(col - 2*cob) < 0.02:
        return 'La columna dice justo el doble de lo que suman los cobros'
    if col > cob:
        return 'Faltan cobros por anotar, o la columna se quedo por encima'
    return 'Los cobros suman mas que la columna: la columna se quedo atras'

wb = Workbook()
# LibreOffice no puede recalcular en este contenedor: se fuerza el recalculo al abrir
# para que Excel rellene todas las formulas la primera vez que se abre el fichero.


# ---------------------------------------------------------------- Detalle
ws = wb.active
ws.title = 'Descuadres por cliente'

ws['A1'] = 'Transportes Araya Franquiz'
ws['A1'].font = Font(name=FUENTE, size=15, bold=True, color=AZUL)
ws['A2'] = 'Facturas donde el importe cobrado de la ficha no cuadra con la suma de sus cobros'
ws['A2'].font = Font(name=FUENTE, size=11, color='FF595959')
ws['A3'] = 'Listado a ' + date(2026,9,22).strftime('%d/%m/%Y') + '  ·  %d facturas de %d clientes' % (len(d), len({x['cliente'] for x in d}))
ws['A3'].font = Font(name=FUENTE, size=10, color='FF808080')

ws['A5'] = ('Para que sirve: en cada una de estas facturas hay dos cifras de lo cobrado que no coinciden. '
            'Una es la casilla "importe cobrado" de la ficha de la factura; la otra es lo que suman los cobros anotados uno a uno. '
            'Hay que decidir cual de las dos es la buena. Escribe tu decision en las dos ultimas columnas, las amarillas.')
ws['A5'].font = Font(name=FUENTE, size=10, italic=True, color='FF595959')
ws['A5'].alignment = Alignment(wrap_text=True, vertical='top')
ws.merge_cells('A5:P5')
ws.row_dimensions[5].height = 40

ws['A6'] = ('Ojo: esto NO infla la deuda del panel. El informe de pendientes se queda con la mayor de las dos cifras, '
            'asi que nadie aparece debiendo de mas por esto. Lo que si hace es enganar a lo que lea la casilla a pelo.')
ws['A6'].font = Font(name=FUENTE, size=10, italic=True, color='FF806000')
ws.merge_cells('A6:P6')

CABS = ['Cliente','CIF','Factura','Fecha','Estado','Estado en Factusol',
        'Total','Importe cobrado (ficha)','Cobros anotados','Diferencia',
        'Nº cobros','Primer cobro','Ultimo cobro','Pendiente en el panel',
        'Que le pasa','Cual es el bueno','Notas']
f0 = 8
for i, t in enumerate(CABS, start=1):
    c = ws.cell(row=f0, column=i, value=t)
    c.font = Font(name=FUENTE, size=10, bold=True, color='FFFFFFFF')
    c.fill = CAB
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    c.border = BORDE
ws.row_dimensions[f0].height = 34

ws.cell(row=f0, column=8).comment = Comment(
    'Es la casilla importe_cobrado de la ficha de la factura.', 'Panel Araya')
ws.cell(row=f0, column=9).comment = Comment(
    'Es la suma de los cobros anotados uno a uno en la factura.', 'Panel Araya')
ws.cell(row=f0, column=14).comment = Comment(
    'Lo que el panel da hoy por pendiente. Se calcula con la mayor de las dos cifras, '
    'por eso el descuadre no infla la deuda.', 'Panel Araya')

fila = f0 + 1
primeras = []   # (cliente, fila_primer_dato, fila_ultimo_dato)
i = 0
while i < len(d):
    cli = d[i]['cliente']
    bloque = [x for x in d if x['cliente'] == cli]
    # cabecera del cliente
    c = ws.cell(row=fila, column=1, value=cli)
    c.font = Font(name=FUENTE, size=11, bold=True, color=AZUL)
    for col in range(1, len(CABS)+1):
        ws.cell(row=fila, column=col).fill = FCLI
        ws.cell(row=fila, column=col).border = BORDE
    ws.cell(row=fila, column=2, value=bloque[0]['cif'] or '')
    ws.cell(row=fila, column=2).font = Font(name=FUENTE, size=10, bold=True, color=AZUL)
    fila += 1
    ini = fila
    for x in bloque:
        ws.cell(row=fila, column=1, value='')
        ws.cell(row=fila, column=2, value='')
        ws.cell(row=fila, column=3, value=x['numero_factura'])
        ws.cell(row=fila, column=4, value=date(*map(int, x['fecha'].split('-'))))
        ws.cell(row=fila, column=5, value=x['estado'].capitalize())
        ws.cell(row=fila, column=6, value=ESTFAC.get(x['estfac_factusol'], str(x['estfac_factusol'])))
        ws.cell(row=fila, column=7, value=x['total'])
        ws.cell(row=fila, column=8, value=x['importe_cobrado'])
        ws.cell(row=fila, column=9, value=x['cobros_anotados'])
        ws.cell(row=fila, column=10, value='=I%d-H%d' % (fila, fila))
        ws.cell(row=fila, column=11, value=x['n_cobros'])
        ws.cell(row=fila, column=12, value=date(*map(int, x['primer_cobro'].split('-'))))
        ws.cell(row=fila, column=13, value=date(*map(int, x['ultimo_cobro'].split('-'))))
        ws.cell(row=fila, column=14, value='=G%d-MAX(H%d,I%d)' % (fila, fila, fila))
        ws.cell(row=fila, column=15, value=diagnostico(x))
        ws.cell(row=fila, column=16, value='')
        ws.cell(row=fila, column=17, value='')
        for col in range(1, len(CABS)+1):
            cc = ws.cell(row=fila, column=col)
            cc.border = BORDE
            cc.font = Font(name=FUENTE, size=10)
        for col in (7,8,9,10,14):
            ws.cell(row=fila, column=col).number_format = EUR
        for col in (4,12,13):
            ws.cell(row=fila, column=col).number_format = FECHA
        ws.cell(row=fila, column=11).alignment = Alignment(horizontal='center')
        ws.cell(row=fila, column=15).alignment = Alignment(wrap_text=True, vertical='top')
        for col in (16,17):
            ws.cell(row=fila, column=col).fill = FEDIT
        fila += 1
    fin = fila - 1
    # subtotal del cliente
    ws.cell(row=fila, column=1, value='')
    ws.cell(row=fila, column=6, value='Suma de %s' % cli)
    ws.cell(row=fila, column=6).alignment = Alignment(horizontal='right')
    for col in (7,8,9,10):
        ws.cell(row=fila, column=col, value='=SUM(%s%d:%s%d)' % (get_column_letter(col), ini, get_column_letter(col), fin))
        ws.cell(row=fila, column=col).number_format = EUR
    for col in range(1, len(CABS)+1):
        cc = ws.cell(row=fila, column=col)
        cc.fill = FSUB
        cc.border = BORDE
        cc.font = Font(name=FUENTE, size=10, bold=True)
    primeras.append((cli, ini, fin))
    fila += 2
    i += len(bloque)

# total general
prim, ult = f0+1, fila-2
ws.cell(row=fila, column=6, value='TOTAL de las %d facturas' % len(d))
ws.cell(row=fila, column=6).alignment = Alignment(horizontal='right')
for col in (7,8,9):
    ws.cell(row=fila, column=col, value='=SUMIF($C$%d:$C$%d,"<>",%s%d:%s%d)' % (prim, ult, get_column_letter(col), prim, get_column_letter(col), ult))
    ws.cell(row=fila, column=col).number_format = EUR
ws.cell(row=fila, column=10, value='=I%d-H%d' % (fila, fila))
ws.cell(row=fila, column=10).number_format = EUR
for col in range(1, len(CABS)+1):
    cc = ws.cell(row=fila, column=col)
    cc.fill = PatternFill('solid', fgColor='FF1F3864')
    cc.font = Font(name=FUENTE, size=11, bold=True, color='FFFFFFFF')
    cc.border = BORDE
fila_total = fila

dv = DataValidation(type='list',
    formula1='"La casilla de la ficha,Los cobros anotados,Hay cobros sin anotar,Es confirming o factoring,Otra cosa - ver notas"',
    allow_blank=True, showDropDown=False)
dv.error = 'Elige una de las opciones de la lista.'
dv.errorTitle = 'Opcion no valida'
ws.add_data_validation(dv)
dv.add('P%d:P%d' % (f0+1, ult))

anchos = {1:42,2:13,3:11,4:11,5:11,6:16,7:14,8:16,9:15,10:13,11:8,12:12,13:12,14:16,15:44,16:26,17:30}
for col, a in anchos.items():
    ws.column_dimensions[get_column_letter(col)].width = a
ws.freeze_panes = 'C%d' % (f0+1)
ws.auto_filter.ref = 'A%d:Q%d' % (f0, ult)
ws.sheet_view.showGridLines = False

# ---------------------------------------------------------------- Resumen
r = wb.create_sheet('Resumen por cliente')
r['A1'] = 'Resumen por cliente'
r['A1'].font = Font(name=FUENTE, size=14, bold=True, color=AZUL)
r['A2'] = 'Cuantas facturas descuadran en cada cliente y por cuanto. Se calcula solo desde la otra hoja.'
r['A2'].font = Font(name=FUENTE, size=10, italic=True, color='FF595959')

CR = ['Cliente','Facturas','Total facturado','Importe cobrado (ficha)','Cobros anotados','Diferencia']
for i, t in enumerate(CR, start=1):
    c = r.cell(row=4, column=i, value=t)
    c.font = Font(name=FUENTE, size=10, bold=True, color='FFFFFFFF')
    c.fill = CAB
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    c.border = BORDE
r.row_dimensions[4].height = 30

HOJA = "'Descuadres por cliente'"
rf = 5
for cli, ini, fin in primeras:
    r.cell(row=rf, column=1, value=cli).font = Font(name=FUENTE, size=10)
    r.cell(row=rf, column=2, value='=COUNTA(%s!$C$%d:$C$%d)' % (HOJA, ini, fin))
    r.cell(row=rf, column=2).alignment = Alignment(horizontal='center')
    for j, col in enumerate(('G','H','I'), start=3):
        r.cell(row=rf, column=j, value='=SUM(%s!$%s$%d:$%s$%d)' % (HOJA, col, ini, col, fin))
        r.cell(row=rf, column=j).number_format = EUR
    r.cell(row=rf, column=6, value='=E%d-D%d' % (rf, rf))
    r.cell(row=rf, column=6).number_format = EUR
    for col in range(1, 7):
        r.cell(row=rf, column=col).border = BORDE
        if r.cell(row=rf, column=col).font.size is None:
            r.cell(row=rf, column=col).font = Font(name=FUENTE, size=10)
    rf += 1

r.cell(row=rf, column=1, value='TOTAL')
for col in range(2, 7):
    r.cell(row=rf, column=col, value='=SUM(%s5:%s%d)' % (get_column_letter(col), get_column_letter(col), rf-1))
    if col > 2:
        r.cell(row=rf, column=col).number_format = EUR
for col in range(1, 7):
    cc = r.cell(row=rf, column=col)
    cc.fill = PatternFill('solid', fgColor='FF1F3864')
    cc.font = Font(name=FUENTE, size=11, bold=True, color='FFFFFFFF')
    cc.border = BORDE
r.cell(row=rf, column=2).alignment = Alignment(horizontal='center')

for col, a in {1:46,2:11,3:18,4:20,5:18,6:16}.items():
    r.column_dimensions[get_column_letter(col)].width = a
r.freeze_panes = 'A5'
r.sheet_view.showGridLines = False

# ---------------------------------------------------------------- Como leerlo
g = wb.create_sheet('Como leerlo')
g['A1'] = 'Como leer este listado'
g['A1'].font = Font(name=FUENTE, size=14, bold=True, color=AZUL)
texto = [
 ('Que es esto', ''),
 ('', 'Cada factura de aqui tiene dos cifras de lo cobrado que no coinciden: la casilla "importe cobrado" de su ficha, y lo que suman los cobros anotados uno a uno. Hay que decidir cual es la buena.'),
 ('', ''),
 ('Donde escribes tu', ''),
 ('', 'En las dos columnas amarillas de la otra hoja: "Cual es el bueno" (con desplegable) y "Notas". El resto no hay que tocarlo.'),
 ('', ''),
 ('Las opciones del desplegable', ''),
 ('La casilla de la ficha', 'La cifra buena es la de la ficha. Entonces faltan cobros por anotar en el panel.'),
 ('Los cobros anotados', 'La cifra buena es la suma de los cobros. Entonces hay que corregir la casilla de la ficha.'),
 ('Hay cobros sin anotar', 'Se cobro de verdad, pero ese cobro no llego nunca al panel. Hay que meterlo con su fecha.'),
 ('Es confirming o factoring', 'El dinero entro por el banco de otra manera y por eso no hay cobro anotado.'),
 ('Otra cosa - ver notas', 'Cualquier otro caso. Explicalo en la columna de Notas.'),
 ('', ''),
 ('Lo que NO pasa', ''),
 ('', 'Esto no hace que nadie aparezca debiendo de mas. El informe de pendientes de cobro se queda siempre con la mayor de las dos cifras, asi que la deuda que ves es la buena. El descuadre solo enganya a lo que lea la casilla directamente, como le paso al expediente de HUBARA el 21/09/2026.'),
 ('', ''),
 ('De donde sale', ''),
 ('', 'De la base del panel, el 22/09/2026: facturas_venta.importe_cobrado frente a la suma de la tabla cobros. Se sacaron las que se diferencian en mas de un centimo y tienen al menos un cobro anotado.'),
]
rr = 3
for tit, cuerpo in texto:
    if tit:
        c = g.cell(row=rr, column=1, value=tit)
        c.font = Font(name=FUENTE, size=11, bold=True, color=AZUL)
    if cuerpo:
        c = g.cell(row=rr, column=2, value=cuerpo)
        c.font = Font(name=FUENTE, size=10)
        c.alignment = Alignment(wrap_text=True, vertical='top')
        g.row_dimensions[rr].height = max(15, 14 * (1 + len(cuerpo)//95))
    rr += 1
g.column_dimensions['A'].width = 30
g.column_dimensions['B'].width = 105
g.sheet_view.showGridLines = False

wb.calculation.fullCalcOnLoad = True
wb.save('/home/user/TRANSPORTES-ARAYA/informes/descuadres_cobrado_20260922.xlsx')
print('guardado; filas de datos:', len(d), 'clientes:', len(primeras), 'fila total:', fila_total)
